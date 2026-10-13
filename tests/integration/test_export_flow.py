import io
from datetime import UTC, date, datetime
from decimal import Decimal
from uuid import UUID

import openpyxl
import pytest
import pytest_asyncio
from httpx import AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession

from src.energetico.domain.models import RelatorioESG
from src.epidemiologico.domain.models import OcorrenciaSanitaria
from src.operacional.domain.models import Insumo, LoteInsumo


@pytest_asyncio.fixture
async def produtor_auth(async_client: AsyncClient) -> tuple[str, UUID, UUID]:
    register_payload = {
        "email": "export_auditor@seed.agro",
        "nome_completo": "Fazenda Export Audit",
        "senha": "SenhaForte123!",
        "role": "ADMIN"
    }
    await async_client.post("/api/v1/auth/registro", json=register_payload)
    login_payload = {"email": "export_auditor@seed.agro", "senha": "SenhaForte123!"}
    resp = await async_client.post("/api/v1/auth/login", json=login_payload)
    token = resp.json()["access_token"]

    headers = {"Authorization": f"Bearer {token}"}
    me_resp = await async_client.get("/api/v1/auth/me", headers=headers)
    user_id = UUID(me_resp.json()["id"])
    prop_id = UUID(me_resp.json()["propriedade_id"])
    return token, user_id, prop_id


@pytest.mark.asyncio
async def test_export_inventario_insumos_xlsx_flow(
    async_client: AsyncClient,
    produtor_auth: tuple[str, UUID, UUID],
    session: AsyncSession
):
    token, _, prop_id = produtor_auth
    headers = {"Authorization": f"Bearer {token}"}

    # Cadastra Insumo e Lote
    insumo = Insumo(
        nome="Fungicida Fox Xpro",
        tipo="FUNGICIDA",
        unidade_medida="L",
        preco_unitario=Decimal("185.50")
    )
    session.add(insumo)
    await session.flush()

    lote = LoteInsumo(
        propriedade_id=prop_id,
        insumo_id=insumo.id,
        codigo_lote="LOTE-TEST-2026",
        quantidade_inicial=Decimal("500.0"),
        quantidade_atual=Decimal("350.0"),
        data_validade=date(2027, 12, 31),
        entrada_em=datetime.now(UTC),
        fornecedor="Bayer CropScience"
    )
    session.add(lote)
    await session.commit()

    # Requisita exportação XLSX
    resp = await async_client.get("/api/v1/operacional/insumos/export/xlsx", headers=headers)
    assert resp.status_code == 200
    expected_xlsx_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert expected_xlsx_type in resp.headers["content-type"]
    assert "inventario_insumos_seed.xlsx" in resp.headers.get("content-disposition", "")

    # Valida bytes do Excel
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert "Inventário de Insumos" in wb.sheetnames
    ws = wb["Inventário de Insumos"]
    assert ws["A1"].value is not None
    assert "SEED AGRO" in str(ws["A1"].value)


@pytest.mark.asyncio
async def test_export_laudo_fitossanitario_pdf_flow(
    async_client: AsyncClient,
    produtor_auth: tuple[str, UUID, UUID],
    session: AsyncSession
):
    token, user_id, prop_id = produtor_auth
    headers = {"Authorization": f"Bearer {token}"}

    # Cadastra Ocorrência Sanitária
    ocorrencia = OcorrenciaSanitaria(
        propriedade_id=prop_id,
        tipo_cultura="Soja",
        categoria="DOENCA",
        agente_identificado="Ferrugem Asiática (Phakopsora pachyrhizi)",
        severidade="CRITICA",
        latitude=-12.5020,
        longitude=-55.7015,
        reportado_por=user_id,
        origem="WEB",
        observado_em=datetime.now(UTC)
    )
    session.add(ocorrencia)
    await session.commit()

    # Requisita exportação PDF
    url_pdf = "/api/v1/epidemiologico/ocorrencias/export/pdf"
    resp = await async_client.get(url_pdf, headers=headers)
    assert resp.status_code == 200
    assert "application/pdf" in resp.headers["content-type"]
    assert "laudo_fitossanitario.pdf" in resp.headers.get("content-disposition", "")
    assert resp.content.startswith(b"%PDF")


@pytest.mark.asyncio
async def test_export_laudo_esg_pdf_flow(
    async_client: AsyncClient,
    produtor_auth: tuple[str, UUID, UUID],
    session: AsyncSession
):
    token, _, prop_id = produtor_auth
    headers = {"Authorization": f"Bearer {token}"}

    relatorio = RelatorioESG(
        propriedade_id=prop_id,
        periodo="2026-09",
        emissao_total_co2e_ton=Decimal("12.450"),
        emissao_por_hectare=Decimal("0.018"),
        pct_energia_renovavel=Decimal("78.5"),
        indicadores_detalhados={"matriz": {"percentual_fossil": 21.5, "breakdown_fontes": {}}},
        status="HOMOLOGADO"
    )
    session.add(relatorio)
    await session.commit()

    # 1. Teste de download por ID específico
    url_id = f"/api/v1/energetico/relatorios/{relatorio.id}/export/pdf"
    resp_id = await async_client.get(url_id, headers=headers)
    assert resp_id.status_code == 200
    assert "application/pdf" in resp_id.headers["content-type"]
    assert resp_id.content.startswith(b"%PDF")

    # 2. Teste de download do último relatório (sem ID específico)
    url_latest = "/api/v1/energetico/relatorios/export/pdf"
    resp_latest = await async_client.get(url_latest, headers=headers)
    assert resp_latest.status_code == 200
    assert "application/pdf" in resp_latest.headers["content-type"]
    assert resp_latest.content.startswith(b"%PDF")
